# -*- coding: utf-8 -*-
"""
Created on Sat Mar 22 07:17:31 2025

@author: FIX
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, Font
from datetime import datetime

# Legge il file Excel
df = pd.read_excel(r"C:\Users\Utente\PycharmProjects\reportManager\assets\posting_TA_25\test.xlsx")
oggi = datetime.today().strftime("%d-%m-%Y")

# Aggiunge un identificativo per i record originali
df['RecordID'] = df.index

# Crea la colonna "Personale" unendo "Preposto" e "Squadra di lavoro"
df['Personale'] = df['Preposto'].fillna('') + ',' + df['Squadra di lavoro'].fillna('')
df['Personale'] = df['Personale'].apply(lambda x: [nome.strip() for nome in x.split(',') if nome.strip() != ''])

# Seleziona e rinomina le colonne:
# - 'Ditta'           -> 'Azienda'
# - 'Descr. PDL'       -> 'Attività'
# - 'Codice EWP'      -> 'Numero EWP'
# - 'Codice PDL'      -> 'Numero PDL'
df_new = df[['Ditta', 'Descr. PDL', 'Codice EWP', 'Codice PDL', 'Personale', 'RecordID']].copy()
df_new = df_new.rename(columns={
    'Ditta': 'Azienda',
    'Descr. PDL': 'Attività',
    'Codice EWP': 'Numero EWP',
    'Codice PDL': 'Numero PDL'
})

# Aggiunge le colonne vuote "Distacchi" e "Tipologia Distacco"
df_new['Distacchi'] = ""
df_new['Tipologia Distacco'] = ""

# Esplode la colonna "Personale" per avere una riga per ciascun nome
df_new = df_new.explode('Personale').reset_index(drop=True)

# Ordine finale delle colonne
final_columns = ['Azienda', 'Attività', 'Personale', 'Numero EWP', 'Numero PDL', 'Distacchi', 'Tipologia Distacco']

# Crea un nuovo Workbook e seleziona il foglio attivo
wb = Workbook()
ws = wb.active

# Definisce lo stile dei bordi sottili
thin_border = Border(left=Side(style='thin'),
                     right=Side(style='thin'),
                     top=Side(style='thin'),
                     bottom=Side(style='thin'))

# Stile per l'intestazione
header_font = Font(bold=True)
header_alignment = Alignment(horizontal="center")

# Colonne per cui mostrare il dato solo nella prima riga del gruppo
cols_single = ['Azienda', 'Attività', 'Numero EWP', 'Numero PDL']

current_row = 1

# Raggruppa per RecordID (ogni gruppo corrisponde a un record originale)
groups = df_new.groupby('RecordID', sort=False)

for record_id, group in groups:
    group = group.reset_index(drop=True)

    # Scrive l'intestazione per il gruppo
    for col_index, header in enumerate(final_columns, start=1):
        cell = ws.cell(row=current_row, column=col_index, value=header)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    current_row += 1

    # Scrive le righe dei dati del gruppo
    for i, row in group.iterrows():
        for col_index, col in enumerate(final_columns, start=1):
            # Nella prima riga del gruppo mostro il dato completo,
            # mentre nelle successive, se la colonna è tra quelle da mostrare una sola volta, lascio la cella vuota
            value = row[col] if i == 0 or col not in cols_single else ""
            cell = ws.cell(row=current_row, column=col_index, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left")
        current_row += 1

    # Inserisce una riga vuota come spaziatura tra i gruppi
    current_row += 1

# Salva il file Excel formattato
output_file = r"C:\Users\Utente\PycharmProjects\reportManager\assets\posting_TA_25\distacchi_TA25_" + oggi + ".xlsx"
wb.save(output_file)
print(f"File Excel generato: {output_file}")
